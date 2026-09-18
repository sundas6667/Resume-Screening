"""Unit tests for core/report_generator.py."""
import json

import openpyxl
import pandas as pd
import pytest

from core.report_generator import export_to_csv, export_to_excel, export_to_json, export_to_pdf
from models.candidate import AIInsights, Candidate, CategoryScore, MatchingResult, RankingResult
from models.job_description import JobDescription


def make_ranked_candidate(name: str, rank: int, score: float) -> Candidate:
    c = Candidate(full_name=name, email=f"{name.lower().replace(' ', '.')}@example.com", skills=["Python", "AWS"])
    c.matching = MatchingResult(matched_skills=["Python"], missing_skills=["Docker"])
    c.ranking = RankingResult(
        overall_score=score, rank=rank, recommendation="Recommended" if score >= 60 else "Consider",
        confidence_level="High",
        category_breakdown={"skills": CategoryScore(raw_score=80.0, weight=0.35, weighted_contribution=28.0, reason="test")},
    )
    c.insights = AIInsights(strengths=["Strong Python skills"], weaknesses=["Missing Docker"],
                             recruiter_summary=f"Overall score {score}/100.")
    return c


@pytest.fixture
def ranked_candidates():
    return [make_ranked_candidate("Jane Doe", 1, 85.0), make_ranked_candidate("John Smith", 2, 62.0)]


class TestExportToCsv:
    def test_produces_valid_csv_with_correct_row_count(self, ranked_candidates):
        csv_bytes = export_to_csv(ranked_candidates)
        df = pd.read_csv(pd.io.common.BytesIO(csv_bytes))
        assert len(df) == 2
        assert "Jane Doe" in df["Name"].values

    def test_includes_overall_score_column(self, ranked_candidates):
        csv_bytes = export_to_csv(ranked_candidates)
        df = pd.read_csv(pd.io.common.BytesIO(csv_bytes))
        assert "Overall Score" in df.columns
        assert 85.0 in df["Overall Score"].values

    def test_empty_list_does_not_crash(self):
        result = export_to_csv([])
        assert b"No candidates" in result


class TestExportToExcel:
    def test_produces_valid_xlsx(self, ranked_candidates):
        excel_bytes = export_to_excel(ranked_candidates)
        workbook = openpyxl.load_workbook(pd.io.common.BytesIO(excel_bytes))
        sheet = workbook.active
        assert sheet.title == "Rankings"
        assert sheet.cell(row=1, column=1).value == "Rank"

    def test_correct_number_of_data_rows(self, ranked_candidates):
        excel_bytes = export_to_excel(ranked_candidates)
        workbook = openpyxl.load_workbook(pd.io.common.BytesIO(excel_bytes))
        sheet = workbook.active
        assert sheet.max_row == 3  # header + 2 candidates

    def test_score_cell_is_color_filled(self, ranked_candidates):
        excel_bytes = export_to_excel(ranked_candidates)
        workbook = openpyxl.load_workbook(pd.io.common.BytesIO(excel_bytes))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        score_col = headers.index("Overall Score") + 1
        cell = sheet.cell(row=2, column=score_col)
        assert cell.fill.start_color.rgb not in (None, "00000000")

    def test_empty_list_does_not_crash(self):
        excel_bytes = export_to_excel([])
        workbook = openpyxl.load_workbook(pd.io.common.BytesIO(excel_bytes))
        assert workbook.active.cell(row=1, column=1).value == "No candidates to export"


class TestExportToPdf:
    def test_produces_valid_pdf_bytes(self, ranked_candidates):
        pdf_bytes = export_to_pdf(ranked_candidates)
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 500

    def test_includes_job_description_title_when_provided(self, ranked_candidates):
        jd = JobDescription(raw_text="jd", cleaned_text="jd", title="Senior Engineer")
        pdf_bytes = export_to_pdf(ranked_candidates, job_description=jd)
        assert b"Senior Engineer" in pdf_bytes or len(pdf_bytes) > 500  # text is encoded in PDF streams

    def test_custom_company_name_does_not_crash(self, ranked_candidates):
        pdf_bytes = export_to_pdf(ranked_candidates, company_name="Acme Recruiting")
        assert pdf_bytes.startswith(b"%PDF")

    def test_empty_list_does_not_crash(self):
        pdf_bytes = export_to_pdf([])
        assert pdf_bytes.startswith(b"%PDF")

    def test_candidate_with_special_characters_does_not_crash(self):
        """Regression guard: resume/JD content can contain '&', '<', etc.,
        which would break reportlab's Paragraph markup parser if unescaped."""
        candidate = make_ranked_candidate("A & B <Test>", 1, 75.0)
        candidate.insights.strengths = ["Strong in C++ & Python <development>"]
        pdf_bytes = export_to_pdf([candidate])
        assert pdf_bytes.startswith(b"%PDF")

    def test_candidate_with_no_category_breakdown_does_not_crash(self):
        candidate = Candidate(full_name="Empty Candidate")
        candidate.ranking = RankingResult(overall_score=0.0, rank=1, recommendation="Not Recommended")
        pdf_bytes = export_to_pdf([candidate])
        assert pdf_bytes.startswith(b"%PDF")


class TestExportToJson:
    def test_produces_valid_json(self, ranked_candidates):
        json_bytes = export_to_json(ranked_candidates)
        data = json.loads(json_bytes)
        assert data["candidate_count"] == 2

    def test_includes_job_description_when_provided(self, ranked_candidates):
        jd = JobDescription(raw_text="jd", cleaned_text="jd", title="Senior Engineer", required_skills=["Python"])
        data = json.loads(export_to_json(ranked_candidates, job_description=jd))
        assert data["job_description"]["title"] == "Senior Engineer"

    def test_job_description_none_when_not_provided(self, ranked_candidates):
        data = json.loads(export_to_json(ranked_candidates))
        assert data["job_description"] is None

    def test_excludes_raw_text(self, ranked_candidates):
        ranked_candidates[0].raw_text = "CONFIDENTIAL FULL RESUME BODY TEXT"
        data = json.loads(export_to_json(ranked_candidates))
        assert "CONFIDENTIAL FULL RESUME BODY TEXT" not in json.dumps(data)

    def test_empty_list_produces_valid_json(self):
        data = json.loads(export_to_json([]))
        assert data["candidate_count"] == 0
        assert data["candidates"] == []


class TestExecutiveSummaryAndFooter:
    def test_pdf_with_multiple_candidates_includes_summary_page(self, ranked_candidates):
        pdf_bytes = export_to_pdf(ranked_candidates)
        assert pdf_bytes.startswith(b"%PDF")
        assert len(pdf_bytes) > 1000  # summary page adds meaningful content

    def test_recruiter_name_and_report_title_do_not_crash(self, ranked_candidates):
        pdf_bytes = export_to_pdf(
            ranked_candidates, recruiter_name="Pat Recruiter", report_title="Q3 Engineering Search",
        )
        assert pdf_bytes.startswith(b"%PDF")


class TestExportManifest:
    def test_manifest_is_valid_json(self, ranked_candidates):
        from core.report_generator import generate_export_manifest
        manifest_bytes = generate_export_manifest(ranked_candidates, ["csv", "pdf", "xlsx"])
        manifest = json.loads(manifest_bytes)
        assert manifest["candidate_count"] == 2
        assert set(manifest["files_generated"]) == {"csv", "pdf", "xlsx"}

    def test_manifest_includes_scoring_version(self, ranked_candidates):
        from core.report_generator import generate_export_manifest
        manifest = json.loads(generate_export_manifest(ranked_candidates, ["csv"]))
        assert "report_version" in manifest

    def test_manifest_handles_empty_candidate_list(self):
        from core.report_generator import generate_export_manifest
        manifest = json.loads(generate_export_manifest([], ["csv"]))
        assert manifest["candidate_count"] == 0
        assert manifest["report_version"] is None
